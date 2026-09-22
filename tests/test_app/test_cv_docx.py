import io

import pytest
from docx import Document

from app.cv_docx import render_docx
from app.cv_markdown import ENTRY_PAGE_BREAK_MARKER, SECTION_PAGE_BREAK_MARKER
from app.errors import ExportError
from domain.models import AssembledCV, ContactItem, Project


def make_cv(**overrides) -> AssembledCV:
    defaults = dict(name="Ada Lovelace", summary="A summary.")
    defaults.update(overrides)
    return AssembledCV(**defaults)


def paragraph_texts(data: bytes) -> list[str]:
    doc = Document(io.BytesIO(data))
    return [p.text for p in doc.paragraphs if p.text.strip()]


def hyperlink_targets(data: bytes) -> set[str]:
    doc = Document(io.BytesIO(data))
    return {rel.target_ref for rel in doc.part.rels.values() if rel.reltype.endswith("hyperlink")}


def test_header_includes_name_and_headline():
    cv = make_cv(headline="Analytical Engineer")

    texts = paragraph_texts(render_docx(cv, sections={}))

    assert texts[0] == "Ada Lovelace"
    assert "Analytical Engineer" in texts


def test_contacts_section_renders_one_bullet_per_contact():
    # Contacts moved out of the header into its own section (Contacts is
    # first in SECTION_TITLES) — this is the same "edited sections dict"
    # path every other section already uses, not header-special-cased.
    cv = make_cv(contacts=[ContactItem(id="c1", label="Email", value="ada@example.com")])
    sections = {"contacts": "- Email: ada@example.com"}

    texts = paragraph_texts(render_docx(cv, sections))

    assert "Email: ada@example.com" in texts


def test_uses_russian_headings_for_a_russian_language_cv():
    # Version 4, Phase 4.3. Unlike plain PDF export (Base-14 Helvetica has
    # no Cyrillic glyphs — see test_cv_pdf.py's equivalent smoke test),
    # DOCX has no font/glyph limitation here: OOXML text is just encoded
    # XML content, so this asserts the actual localized heading text
    # extracts correctly, not just "doesn't crash" — confirmed live.
    cv = make_cv(name="Иван Иванов", language="ru", summary="О себе.")
    sections = {"summary": "Текст на русском."}

    texts = paragraph_texts(render_docx(cv, sections))

    assert "Иван Иванов" in texts
    assert "О себе" in texts
    assert "Summary" not in texts


def test_only_non_blank_sections_get_a_heading():
    cv = make_cv()
    sections = {"summary": "Some text.", "skills": "", "experience": "   "}

    texts = paragraph_texts(render_docx(cv, sections))

    assert "Summary" in texts
    assert "Skills" not in texts
    assert "Experience" not in texts


def test_section_heading_is_kept_with_the_following_paragraph():
    # Real bug found from a real export: without this, Word is free to
    # break the page right after a heading, stranding it alone with its
    # content starting fresh on the next page.
    cv = make_cv()
    sections = {"summary": "Some text."}

    doc = Document(io.BytesIO(render_docx(cv, sections)))
    heading = next(p for p in doc.paragraphs if p.text == "Summary")

    assert heading.paragraph_format.keep_with_next is True


def test_inline_bold_span_becomes_a_bold_run():
    cv = make_cv()
    sections = {"experience": "**Engineer** — Acme"}

    doc = Document(io.BytesIO(render_docx(cv, sections)))
    line = next(p for p in doc.paragraphs if p.text == "Engineer — Acme")

    assert line.runs[0].text == "Engineer"
    assert line.runs[0].bold is True
    assert line.runs[1].bold is not True


def test_inline_underline_span_becomes_an_underlined_run():
    # Phase 25: __underline__ is a new markdown_inline.py span, only ever
    # produced by cv_markdown.py::_serialize_runs for a block with
    # explicit runs — exercised end to end here via the same plain
    # sections dict render_docx already accepts.
    cv = make_cv()
    sections = {"experience": "__Engineer__ — Acme"}

    doc = Document(io.BytesIO(render_docx(cv, sections)))
    line = next(p for p in doc.paragraphs if p.text == "Engineer — Acme")

    assert line.runs[0].text == "Engineer"
    assert line.runs[0].underline is True
    assert line.runs[1].underline is not True


def test_bullet_line_uses_list_bullet_style():
    cv = make_cv()
    sections = {"skills": "- Python\n- SQL"}

    doc = Document(io.BytesIO(render_docx(cv, sections)))
    bullet_paragraphs = [p for p in doc.paragraphs if p.style.name == "List Bullet"]

    assert [p.text for p in bullet_paragraphs] == ["Python", "SQL"]


def test_contact_url_becomes_a_clickable_hyperlink():
    cv = make_cv()
    sections = {"contacts": "- Email: ada@example.com\n- Portfolio: https://ada.example.com/portfolio"}

    data = render_docx(cv, sections)

    assert hyperlink_targets(data) == {"https://ada.example.com/portfolio"}


def test_project_url_becomes_a_clickable_hyperlink():
    cv = make_cv(projects=[Project(id="p1", name="Analytical Engine", url="https://example.com/engine")])
    sections = {"projects": "- **Analytical Engine** (https://example.com/engine)"}

    data = render_docx(cv, sections)

    assert hyperlink_targets(data) == {"https://example.com/engine"}


def test_portfolio_link_url_becomes_a_clickable_hyperlink():
    cv = make_cv()
    sections = {"portfolio_links": "- Concept art — https://example.com/portfolio"}

    data = render_docx(cv, sections)

    assert hyperlink_targets(data) == {"https://example.com/portfolio"}


def test_publication_url_becomes_a_clickable_hyperlink():
    cv = make_cv()
    sections = {"publications": "- **Scaling LiveOps** — Medium (https://example.com/article)"}

    data = render_docx(cv, sections)

    assert hyperlink_targets(data) == {"https://example.com/article"}


# ----- Phase 30: manual pagination control (plain DOCX) -------------------


def test_leading_section_marker_sets_the_headings_page_break_flag_and_is_not_rendered():
    cv = make_cv()
    sections = {"summary": f"{SECTION_PAGE_BREAK_MARKER}\nSome text."}

    doc = Document(io.BytesIO(render_docx(cv, sections)))
    heading = next(p for p in doc.paragraphs if p.text == "Summary")

    assert heading.paragraph_format.page_break_before is True
    assert not any(SECTION_PAGE_BREAK_MARKER in p.text for p in doc.paragraphs)


def test_section_without_a_leading_marker_leaves_the_heading_flag_unset():
    cv = make_cv()
    sections = {"summary": "Some text."}

    doc = Document(io.BytesIO(render_docx(cv, sections)))
    heading = next(p for p in doc.paragraphs if p.text == "Summary")

    assert not heading.paragraph_format.page_break_before


def test_mid_body_marker_sets_the_page_break_flag_on_the_paragraph_that_follows_it():
    cv = make_cv()
    sections = {"skills": f"- Mentoring\n{ENTRY_PAGE_BREAK_MARKER}\n- Roadmapping"}

    doc = Document(io.BytesIO(render_docx(cv, sections)))
    first = next(p for p in doc.paragraphs if p.text == "Mentoring")
    second = next(p for p in doc.paragraphs if p.text == "Roadmapping")

    assert not first.paragraph_format.page_break_before
    assert second.paragraph_format.page_break_before is True
    assert not any(ENTRY_PAGE_BREAK_MARKER in p.text for p in doc.paragraphs)


def test_entry_marker_on_the_bodys_first_line_leaves_the_heading_flag_unset():
    # Distinct from a section-level marker even though both would
    # otherwise look identical as body's leading line: this one came from
    # the section's *first entry* being flagged, not the section itself
    # (see cv_markdown.py's ENTRY_PAGE_BREAK_MARKER docstring) — the
    # heading must NOT jump to the new page, only the entry.
    cv = make_cv()
    sections = {"skills": f"{ENTRY_PAGE_BREAK_MARKER}\n- Roadmapping"}

    doc = Document(io.BytesIO(render_docx(cv, sections)))
    heading = next(p for p in doc.paragraphs if p.text == "Skills")
    entry = next(p for p in doc.paragraphs if p.text == "Roadmapping")

    assert not heading.paragraph_format.page_break_before
    assert entry.paragraph_format.page_break_before is True


def test_raises_export_error_when_content_exceeds_size_limit(monkeypatch):
    import app.cv_docx as cv_docx_module

    monkeypatch.setattr(cv_docx_module, "_MAX_EXPORT_BYTES", 10)
    cv = make_cv()

    with pytest.raises(ExportError, match="2MB"):
        render_docx(cv, sections={"summary": "Some text."})
