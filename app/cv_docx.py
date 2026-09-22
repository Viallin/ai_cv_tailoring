"""Renders an AssembledCV + edited section text into a .docx document.

Mirrors app/cv_markdown.py's header/section traversal, but builds on the
*edited* section text (see ui/main_window.py's export dispatch) rather than
re-rendering from the AssembledCV — per Phase 6/8's "edited text is never
re-parsed back into the domain model" rule, export must show exactly what
the user last edited. The header is the exception: it isn't user-editable
in the UI, so it's built straight from `cv` rather than parsed back out of
Markdown text.

`python-docx` is imported lazily (mirrors providers/gemini_provider.py's
and app/resume_reader.py's lazy-import convention) since it's only needed
at actual export time.
"""

from __future__ import annotations

import io

from app.cv_locales import section_titles_for
from app.cv_markdown import (
    ENTRY_PAGE_BREAK_MARKER,
    RunsRenderable,
    SubheadingGroupMember,
    group_entries_by_subheading,
    split_leading_page_break_marker,
)
from app.cv_templates import DocxTemplate, resolve_docx_template
from app.errors import ExportError
from app.markdown_inline import parse_inline_runs
from domain.models import AssembledCV, PrintDocument

_MAX_EXPORT_BYTES = 2_000_000


def render_docx(cv: AssembledCV, sections: dict[str, str]) -> bytes:
    try:
        from docx import Document
    except ImportError as exc:
        raise ExportError(
            "DOCX export requires the 'python-docx' package (listed in pyproject.toml)."
        ) from exc

    try:
        document = Document()
        _render_header(document, cv)
        for key, title in section_titles_for(cv.language).items():
            body = sections.get(key, "")
            if not body.strip():
                continue
            # Phase 30 — manual pagination control's plain-export
            # equivalent (see SECTION_PAGE_BREAK_MARKER's docstring). A
            # section-level break forces the heading itself onto a new
            # page, so it's read off `body`'s own leading line here,
            # before the heading paragraph even exists.
            section_break, body = split_leading_page_break_marker(body)
            heading = document.add_heading(title, level=2)
            # Without this, Word is free to break the page right after a
            # heading, stranding it alone with its actual content starting
            # fresh on the next page — a real bug found from a real
            # export. Mirrors the `break-after: avoid-page` fix already
            # applied to the on-screen/Playwright-PDF path's CSS
            # (frontend/src/index.css, Phase 16c) for the exact same
            # symptom; python-docx has no CSS, so `keep_with_next` on the
            # heading paragraph is the direct equivalent — "keep me on the
            # same page as whatever paragraph follows."
            heading.paragraph_format.keep_with_next = True
            heading.paragraph_format.page_break_before = section_break
            _render_body(document, body)

        buffer = io.BytesIO()
        document.save(buffer)
        data = buffer.getvalue()
    except Exception as exc:  # noqa: BLE001 - translate any python-docx error
        raise ExportError(f"Could not build DOCX export: {exc}") from exc

    if len(data) > _MAX_EXPORT_BYTES:
        raise ExportError(
            f"Generated DOCX is {len(data) / 1_000_000:.1f}MB, exceeding the 2MB export limit."
        )
    return data


def render_templated_docx(cv: AssembledCV, document: PrintDocument, template_id: str) -> bytes:
    """Phase 16b — the template-aware DOCX export. Walks `document`, the
    *edited* structured document (Phase 16a's A4 Preview tab state: already
    reordered/toggled by the user), rather than the flat `sections` dict
    `render_docx` uses — this is what makes templated DOCX respect the
    Preview tab's edits instead of silently reverting to unedited content,
    matching what the templated PDF export (app/cv_pdf_playwright.py) does.

    Deliberately doesn't reuse `render_docx`'s "List Bullet" style for
    bullets (Word's native numbering, not a literal character) — a
    template's chosen bullet marker can't be expressed through that
    without `numbering.xml` surgery, so this path uses literal-character-
    prefixed plain paragraphs instead, mirroring app/cv_pdf.py's own
    already-solved version of the same problem. `render_docx`'s own
    bullet handling is untouched.
    """
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise ExportError(
            "DOCX export requires the 'python-docx' package (listed in pyproject.toml)."
        ) from exc

    # Version 4, Phase 4.4 — Calibri substitution for any non-English
    # profile (Georgia/Geist's Cyrillic coverage can't be verified from
    # Python, and Word's own silent-substitution behavior isn't something
    # to gamble on) — see resolve_docx_template's own docstring.
    template = resolve_docx_template(template_id, cv.language)

    try:
        docx_document = Document()
        _render_header(docx_document, cv, heading_font=template.heading_font)
        for section in document.sections:
            if not section.included:
                continue
            entries = [entry for entry in section.entries if entry.included]
            if not entries:
                continue
            # Post-31 — see app/cv_markdown.py's identical check (and
            # domain/models.py's PrintDocumentBlock.kind docstring) for
            # the full rationale: an already-compact Skills/Technologies
            # section (structuredDocument.ts's buildCompactSkillEntries)
            # must never re-enter group_entries_by_subheading below —
            # grouping two already-final lines together would wrongly
            # fuse them onto one.
            is_compact_skills_section = section.key in ("skills", "technologies") and all(
                entry.kind == "compact" for entry in entries
            )
            heading = docx_document.add_heading(section.title, level=2)
            heading.paragraph_format.space_before = Pt(template.section_spacing_pt)
            # See render_docx's identical line for why — same real bug,
            # same fix, same CSS precedent (Phase 16c) it mirrors.
            heading.paragraph_format.keep_with_next = True
            # Phase 30 — manual pagination control. Word composes this
            # with `keep_with_next` on its own: nothing here needs to
            # special-case a section that happens to be first in the
            # document — same as CSS `break-before: page`
            # (PrintSectionBlock.tsx/PrintExperienceSection.tsx), forcing
            # a break here only pushes the CV header (name/headline) alone
            # onto its own page, exactly the same outcome the on-screen
            # guide and templated-PDF export already produce for it.
            heading.paragraph_format.page_break_before = section.page_break_before
            _style_runs(heading, template.heading_font)
            if section.key == "experience":
                for entry in entries:
                    # The role/position header — bold, no bullet marker
                    # (italic instead for an is_gap entry), matching
                    # app/cv_markdown.py's render_experience_section
                    # exactly. Its bullets are the ones that get a
                    # marker; grouped-by-project bullets (kind ==
                    # "subheading") render as an italic sub-line instead,
                    # matching _render_bullets_grouped_by_project.
                    #
                    # Phase 25: an explicit `entry.runs`/`bullet.runs`
                    # (Experience is the one branch here that walks real
                    # PrintDocumentEntry/PrintDocumentBlock objects
                    # directly, not group_entries_by_subheading's
                    # text-only groups — see that helper's docstring for
                    # why every other section stays text-only this
                    # phase) takes over from the bold/italic default
                    # entirely, same "runs is authoritative" rule as the
                    # plain-export bridge (cv_markdown.py::
                    # _block_markdown_text).
                    _add_paragraph_with_runs(
                        docx_document,
                        entry,
                        template,
                        bold=not entry.locked,
                        italic=bool(entry.locked),
                        page_break_before=entry.page_break_before,
                    )
                    for bullet in entry.bullets or []:
                        if not bullet.included:
                            continue
                        if bullet.kind == "subheading":
                            _add_paragraph_with_runs(docx_document, bullet, template, italic=True)
                        else:
                            _add_paragraph_with_runs(docx_document, bullet, template, bullet=True)
            elif is_compact_skills_section:
                # Post-31 — each entry is already a whole pre-joined,
                # final line (structuredDocument.ts's
                # buildCompactSkillEntries — see domain/models.py's
                # PrintDocumentBlock.kind docstring): no bullet character
                # (mirrors PrintSectionBlock.tsx's identical no-bulletChar
                # treatment), rendered via `_add_paragraph_with_runs` so a
                # category entry's own bold `runs` (its "Category: "
                # label) paints, same as Experience gets just above.
                for entry in entries:
                    _add_paragraph_with_runs(docx_document, entry, template, page_break_before=entry.page_break_before)
            else:
                # kind == "subheading" entries currently only appear for
                # Skills/Technologies category names (see
                # structuredDocument.ts's buildCategoryGroupedSection) — a
                # header and the members below it (until the next header)
                # render as one compact bold-prefixed paragraph instead of
                # a heading paragraph + one bullet per member, matching
                # app/cv_markdown.py's shared `group_entries_by_subheading`
                # (also used by render_sections_from_document there) and
                # PrintSectionBlock.tsx's identical grouping (so the
                # templated DOCX and the Playwright-screenshotted
                # templated PDF stay visually consistent with each
                # other). Every other section's entries have no
                # subheadings at all, so this is a no-op reshuffle for
                # them — still one bulleted paragraph per entry, matching
                # cv_markdown.py's render_education_section/
                # render_languages_section/etc. Only ever reached for a
                # *legacy*, not-yet-migrated Skills/Technologies section —
                # `is_compact_skills_section` above handles an already-
                # migrated one instead, never entering
                # group_entries_by_subheading at all (grouping two
                # already-compact lines together would wrongly fuse them
                # onto one).
                #
                # Post-30 polish — two more deliberate exceptions to that
                # "one bulleted paragraph per member" default, both
                # reported directly and both mirrored exactly in
                # PrintSectionBlock.tsx's own `renderFlatGroup` (see its
                # docstring) so the templated DOCX stays visually
                # consistent with the Playwright-screenshotted templated
                # PDF for the same document:
                # - Summary isn't a list — no bullet character, one plain
                #   paragraph per member (still one paragraph each, not
                #   fused, so multiple Summary paragraphs stay distinct).
                # - Skills/Technologies — a long flat skill list read as
                #   too many separate lines; every member in the group
                #   joins onto one `" · "`-separated paragraph instead,
                #   the same "compact single line" treatment a
                #   *categorized* group already gets just above, minus
                #   the category label. Only the first member's own
                #   `page_break_before` has anywhere left to apply once
                #   every member is fused onto one paragraph — mirrors a
                #   category's own flag.
                for group in group_entries_by_subheading(entries):
                    if group.category is not None:
                        _add_grouped_category_paragraph(
                            docx_document,
                            group.category,
                            group.members,
                            template,
                            page_break_before=group.category_page_break_before,
                        )
                    elif section.key == "summary":
                        # Phase 31 fix, found during implementation, not
                        # anticipated in planning: despite Phase 25's own
                        # claim of "Summary + Experience" runs support in
                        # every export path, this branch was still
                        # `.text`-only — `_add_paragraph_with_runs` (the
                        # Experience branch's own helper, a couple of
                        # branches above) takes a full block and already
                        # falls back to identical plain-paragraph,
                        # no-bullet behavior when `runs` is unset, so this
                        # is a drop-in swap, not a new code path.
                        for member in group.members:
                            _add_paragraph_with_runs(
                                docx_document,
                                member,
                                template,
                                page_break_before=member.page_break_before,
                            )
                    elif section.key in ("skills", "technologies") and group.members:
                        _add_styled_paragraph(
                            docx_document,
                            " · ".join(member.text for member in group.members),
                            template,
                            page_break_before=group.members[0].page_break_before,
                        )
                    else:
                        # Phase 31 — the generic "one bulleted paragraph
                        # per member" path for every section without a
                        # more specific case above (Contacts, Education,
                        # Certifications, Awards, Publications, Volunteer
                        # Experience, Portfolio Links). `_add_paragraph_
                        # with_runs` (not `_add_styled_paragraph`, which
                        # only ever takes a bare string) honors an
                        # explicit `member.runs`/`member.alignment` when
                        # the user formatted it via the BubbleMenu,
                        # falling back to the exact same plain-bulleted
                        # behavior `_add_styled_paragraph` used to provide
                        # otherwise — same call already used for
                        # Experience/compact-Skills above.
                        for member in group.members:
                            _add_paragraph_with_runs(
                                docx_document,
                                member,
                                template,
                                bullet=True,
                                page_break_before=member.page_break_before,
                            )

        buffer = io.BytesIO()
        docx_document.save(buffer)
        data = buffer.getvalue()
    except Exception as exc:  # noqa: BLE001 - translate any python-docx error
        raise ExportError(f"Could not build templated DOCX export: {exc}") from exc

    if len(data) > _MAX_EXPORT_BYTES:
        raise ExportError(
            f"Generated DOCX is {len(data) / 1_000_000:.1f}MB, exceeding the 2MB export limit."
        )
    return data


def _add_grouped_category_paragraph(
    document,
    category: str,
    members: list[SubheadingGroupMember],
    template: DocxTemplate,
    *,
    page_break_before: bool = False,
) -> None:
    """One paragraph for a whole category group: a bold run for
    `category`, then (if it has any members) a plain run for
    `": " + "; ".join(member.text for member in members)`. A category with
    zero members (an edge case — an empty category in the edited
    document) still renders as just the bold label, no trailing colon.
    `page_break_before` is the source subheading entry's own Phase 30
    flag — see `render_templated_docx`'s call site."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.page_break_before = page_break_before
    bold_run = paragraph.add_run(category)
    bold_run.font.name = template.body_font
    bold_run.bold = True
    if members:
        run = paragraph.add_run(f": {'; '.join(member.text for member in members)}")
        run.font.name = template.body_font


def _add_styled_paragraph(
    document,
    text: str,
    template: DocxTemplate,
    *,
    bold: bool = False,
    italic: bool = False,
    bullet: bool = False,
    page_break_before: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.page_break_before = page_break_before
    prefix = f"{template.bullet_char} " if bullet else ""
    run = paragraph.add_run(f"{prefix}{text}")
    run.font.name = template.body_font
    run.bold = bold
    run.italic = italic


def _add_paragraph_with_runs(
    document,
    block: RunsRenderable,
    template: DocxTemplate,
    *,
    bold: bool = False,
    italic: bool = False,
    bullet: bool = False,
    page_break_before: bool = False,
) -> None:
    """Phase 25's runs-aware sibling of `_add_styled_paragraph`: when
    `block.runs` is set, renders each run as its own docx run — bold/
    italic/underline/hyperlink, full multi-attribute fidelity per run,
    unlike the plain-export markdown bridge (cv_markdown.py::
    _serialize_runs), which never goes through text at all here — instead
    of the single bold/italic run `_add_styled_paragraph` always builds
    from `block.text`. Falls back to that exact single-run behavior when
    `runs` is unset. `block.alignment`, when set, applies either way —
    it's a paragraph attribute, not a run attribute.

    Phase 25 scoped every call site to Experience only; Phase 31 widened
    that to every section `render_templated_docx` renders (Summary, the
    generic per-section-member branch, and the already-runs-aware
    compact-Skills/Technologies branch) — `block` accepts `RunsRenderable`
    (app/cv_markdown.py), not just a full `PrintDocumentBlock`, so a
    `SubheadingGroupMember` (the lighter carrier
    `group_entries_by_subheading`'s ungrouped branch produces) works here
    too, not just a real `PrintDocumentEntry`.

    `page_break_before` (Phase 30) is only ever passed for a whole entry's
    own paragraph, never a bullet — for Experience specifically that's the
    role-header paragraph (see `render_templated_docx`'s call site; one
    paragraph starting the page is enough to carry the whole entry,
    bullets included, matching PrintExperienceSection.tsx's identical
    scoping), and for every other section it's that section's own
    standalone member paragraph — bullets never carry the flag at all
    (`PrintDocumentBlock` doesn't have the field; only
    `PrintDocumentEntry`/`PrintDocumentSection` do).
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.page_break_before = page_break_before
    prefix = f"{template.bullet_char} " if bullet else ""

    if block.runs:
        if prefix:
            prefix_run = paragraph.add_run(prefix)
            prefix_run.font.name = template.body_font
        for run in block.runs:
            if run.link:
                _add_hyperlink(paragraph, run.link, run.text, bold=run.bold, italic=run.italic)
            else:
                docx_run = paragraph.add_run(run.text)
                docx_run.font.name = template.body_font
                docx_run.bold = run.bold
                docx_run.italic = run.italic
                docx_run.underline = run.underline
    else:
        run = paragraph.add_run(f"{prefix}{block.text}")
        run.font.name = template.body_font
        run.bold = bold
        run.italic = italic

    _ALIGNMENTS = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
    if block.alignment in _ALIGNMENTS:
        paragraph.alignment = _ALIGNMENTS[block.alignment]


def _style_runs(paragraph, font_name: str) -> None:
    for run in paragraph.runs:
        run.font.name = font_name


def _render_header(document, cv: AssembledCV, heading_font: str | None = None) -> None:
    # Contacts moved out of the header into its own section (Contacts is
    # now first in SECTION_TITLES) so each contact can be independently
    # included/excluded like every other CV fact — see
    # app/cv_markdown.py::render_contacts_section.
    heading = document.add_heading(cv.name, level=1)
    if heading_font:
        _style_runs(heading, heading_font)
    if cv.headline:
        run = document.add_paragraph().add_run(cv.headline)
        run.italic = True


def _render_body(document, body: str) -> None:
    # Phase 30 — an entry/category-group's own forced break, encoded as
    # its own marker-only line by render_sections_from_document (see
    # ENTRY_PAGE_BREAK_MARKER's docstring). Deferred to the next real
    # paragraph rather than spent on an empty placeholder paragraph of its
    # own — a paragraph's `page_break_before` is a property of that
    # paragraph itself, the same "set it on the content that should start
    # the page" shape render_templated_docx's own Phase 30 handling uses.
    pending_page_break = False
    for line in body.split("\n"):
        if line == ENTRY_PAGE_BREAK_MARKER:
            pending_page_break = True
            continue
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            text = stripped[2:]
        else:
            paragraph = document.add_paragraph()
            text = stripped
        if pending_page_break:
            paragraph.paragraph_format.page_break_before = True
            pending_page_break = False
        _add_runs(paragraph, text)


def _add_runs(paragraph, text: str) -> None:
    for run in parse_inline_runs(text):
        if run.url:
            _add_hyperlink(paragraph, run.url, run.text, bold=run.bold, italic=run.italic)
        else:
            docx_run = paragraph.add_run(run.text)
            docx_run.bold = run.bold
            docx_run.italic = run.italic
            docx_run.underline = run.underline


def _add_hyperlink(paragraph, url: str, text: str, bold: bool, italic: bool) -> None:
    """Adds a real clickable hyperlink run — python-docx has no high-level
    API for this, so it's built via the underlying OOXML directly. Standard
    workaround, e.g. python-docx issue #384."""
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run_elem = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")  # standard Office hyperlink blue
    run_props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.append(underline)
    if bold:
        run_props.append(OxmlElement("w:b"))
    if italic:
        run_props.append(OxmlElement("w:i"))

    run_elem.append(run_props)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    run_elem.append(text_elem)
    hyperlink.append(run_elem)
    paragraph._p.append(hyperlink)
